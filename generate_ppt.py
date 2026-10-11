"""
Generate Review 1 Presentation — Library Circulation & Digital Resource Management System
Premium dark-themed PowerPoint with professional styling.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ── Color Palette ──
BG_DARK      = RGBColor(0x0F, 0x11, 0x1A)   # Deep navy-black
BG_CARD      = RGBColor(0x1A, 0x1D, 0x2E)   # Card background
ACCENT_BLUE  = RGBColor(0x3B, 0x82, 0xF6)   # Vibrant blue
ACCENT_CYAN  = RGBColor(0x06, 0xB6, 0xD4)   # Cyan
ACCENT_GREEN = RGBColor(0x10, 0xB9, 0x81)   # Emerald
ACCENT_PURPLE= RGBColor(0x8B, 0x5C, 0xF6)   # Purple
ACCENT_AMBER = RGBColor(0xF5, 0x9E, 0x0B)   # Amber
ACCENT_ROSE  = RGBColor(0xF4, 0x3F, 0x5E)   # Rose
WHITE        = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY   = RGBColor(0xCB, 0xD5, 0xE1)   # Muted text
MID_GRAY     = RGBColor(0x94, 0xA3, 0xB8)   # Secondary text
BORDER_COLOR = RGBColor(0x33, 0x3A, 0x50)   # Subtle borders

# ── Helpers ──
def set_slide_bg(slide, color=BG_DARK):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_shape_rect(slide, left, top, width, height, fill_color, border_color=None, corner_radius=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if corner_radius else MSO_SHAPE.RECTANGLE,
        left, top, width, height
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

def set_text(shape, text, font_size=14, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT, font_name='Calibri'):
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return tf

def add_text_box(slide, left, top, width, height, text, font_size=14, color=WHITE, bold=False, alignment=PP_ALIGN.LEFT, font_name='Calibri'):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return tf

def add_bullet_text(tf, text, font_size=13, color=LIGHT_GRAY, bold=False, indent_level=0, space_before=Pt(4), font_name='Calibri'):
    p = tf.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.level = indent_level
    p.space_before = space_before
    return p

def add_accent_bar(slide, left, top, width, height, color):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    return bar

def add_section_number(slide, number, accent_color=ACCENT_BLUE):
    """Add a floating section number badge"""
    badge = add_shape_rect(slide, Inches(0.6), Inches(0.5), Inches(0.7), Inches(0.7), accent_color, corner_radius=True)
    set_text(badge, str(number), font_size=28, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    badge.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    badge.text_frame.word_wrap = False

def add_slide_title(slide, title, subtitle=None, accent_color=ACCENT_BLUE):
    """Standard slide title with accent bar"""
    add_accent_bar(slide, Inches(0.6), Inches(1.4), Inches(0.06), Inches(0.5), accent_color)
    add_text_box(slide, Inches(0.85), Inches(1.25), Inches(10), Inches(0.6), title,
                 font_size=28, color=WHITE, bold=True)
    if subtitle:
        add_text_box(slide, Inches(0.85), Inches(1.85), Inches(10), Inches(0.4), subtitle,
                     font_size=14, color=MID_GRAY)

def make_card(slide, left, top, width, height, title, items, accent_color=ACCENT_BLUE, title_size=14, item_size=12):
    """Reusable info card"""
    card = add_shape_rect(slide, left, top, width, height, BG_CARD, BORDER_COLOR, corner_radius=True)
    add_accent_bar(slide, left, top, width, Inches(0.05), accent_color)

    tf = add_text_box(slide, left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), Inches(0.35),
                      title, font_size=title_size, color=accent_color, bold=True)

    content_tf = add_text_box(slide, left + Inches(0.2), top + Inches(0.5),
                              width - Inches(0.4), height - Inches(0.6),
                              "", font_size=item_size, color=LIGHT_GRAY)
    content_tf.paragraphs[0].text = ""
    first = True
    for item in items:
        if first:
            content_tf.paragraphs[0].text = item
            content_tf.paragraphs[0].font.size = Pt(item_size)
            content_tf.paragraphs[0].font.color.rgb = LIGHT_GRAY
            content_tf.paragraphs[0].font.name = 'Calibri'
            content_tf.paragraphs[0].space_before = Pt(2)
            first = False
        else:
            add_bullet_text(content_tf, item, font_size=item_size, color=LIGHT_GRAY, space_before=Pt(2))

# ═══════════════════════════════════════════════════════
#  SLIDE 1 — TITLE SLIDE
# ═══════════════════════════════════════════════════════

slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
set_slide_bg(slide)

# Accent shapes — decorative
add_shape_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.08), ACCENT_BLUE)

# Big accent circle (decorative)
circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.5), Inches(-1.5), Inches(5), Inches(5))
circle.fill.solid()
circle.fill.fore_color.rgb = RGBColor(0x1E, 0x29, 0x3B)
circle.line.fill.background()

circle2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.5), Inches(4.5), Inches(4), Inches(4))
circle2.fill.solid()
circle2.fill.fore_color.rgb = RGBColor(0x17, 0x20, 0x33)
circle2.line.fill.background()

# Title text
add_text_box(slide, Inches(0.8), Inches(1.5), Inches(9), Inches(0.5),
             "DESIGN AND IMPLEMENTATION OF A DATABASE MANAGEMENT SYSTEM FOR",
             font_size=15, color=ACCENT_CYAN, bold=True)

add_text_box(slide, Inches(0.8), Inches(2.1), Inches(9), Inches(1.2),
             "Library Circulation &\nDigital Resource Management",
             font_size=40, color=WHITE, bold=True)

add_accent_bar(slide, Inches(0.8), Inches(3.5), Inches(2.5), Inches(0.06), ACCENT_BLUE)

add_text_box(slide, Inches(0.8), Inches(3.8), Inches(5), Inches(0.4),
             "REVIEW 1  |  Project Based Learning",
             font_size=16, color=MID_GRAY, bold=False)

add_text_box(slide, Inches(0.8), Inches(4.5), Inches(5), Inches(0.4),
             "Database Management Systems",
             font_size=14, color=RGBColor(0x64, 0x74, 0x8B))

# Bottom bar
add_shape_rect(slide, Inches(0), Inches(7.0), Inches(13.333), Inches(0.5), RGBColor(0x14, 0x17, 0x25))
add_text_box(slide, Inches(0.8), Inches(7.05), Inches(6), Inches(0.4),
             "September 2026", font_size=11, color=MID_GRAY)

# ═══════════════════════════════════════════════════════
#  SLIDE 2 — AGENDA / TABLE OF CONTENTS
# ═══════════════════════════════════════════════════════

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_shape_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT_BLUE)

add_text_box(slide, Inches(0.8), Inches(0.5), Inches(5), Inches(0.5),
             "AGENDA", font_size=32, color=WHITE, bold=True)
add_accent_bar(slide, Inches(0.8), Inches(1.1), Inches(1.5), Inches(0.05), ACCENT_BLUE)

agenda_items = [
    ("01", "Problem Identification", ACCENT_BLUE),
    ("02", "Scope & Objectives", ACCENT_CYAN),
    ("03", "Users & User Roles", ACCENT_GREEN),
    ("04", "Functional Requirements", ACCENT_PURPLE),
    ("05", "Non-Functional Requirements & Business Rules", ACCENT_AMBER),
    ("06", "ER Diagram", ACCENT_ROSE),
    ("07", "Relational Schema", ACCENT_BLUE),
]

for i, (num, title, color) in enumerate(agenda_items):
    y = Inches(1.6 + i * 0.72)
    # Number badge
    badge = add_shape_rect(slide, Inches(1.0), y, Inches(0.55), Inches(0.5), color, corner_radius=True)
    set_text(badge, num, font_size=16, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    # Title
    add_text_box(slide, Inches(1.8), y + Inches(0.05), Inches(8), Inches(0.4),
                 title, font_size=18, color=WHITE, bold=False)
    # Line
    add_accent_bar(slide, Inches(1.8), y + Inches(0.5), Inches(8), Inches(0.01), BORDER_COLOR)

# ═══════════════════════════════════════════════════════
#  SLIDE 3 — PROBLEM IDENTIFICATION
# ═══════════════════════════════════════════════════════

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_shape_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT_BLUE)

add_section_number(slide, 1, ACCENT_BLUE)
add_slide_title(slide, "Problem Identification",
                "Key challenges with manual and legacy library management systems")

problems = [
    ("Unreliable Tracking", "Manual ledgers lead to discrepancies in book locations, availability, and member histories.", ACCENT_BLUE),
    ("No Digital Resource Mgmt", "Older systems lack infrastructure to track e-books, journals, audiobooks, and digital access.", ACCENT_CYAN),
    ("Manual Fine Calculation", "Prone to human error, inconsistencies, and disputes; creates unnecessary admin overhead.", ACCENT_GREEN),
    ("Inefficient Reservations", "No automated queue system leads to chaotic, unfair distribution of popular titles.", ACCENT_PURPLE),
    ("Data Redundancy", "Duplicated information across registers causes anomalies and unreliable reporting.", ACCENT_AMBER),
]

for i, (title, desc, color) in enumerate(problems):
    col = i % 3
    row = i // 3
    left = Inches(0.6 + col * 4.1)
    top = Inches(2.5 + row * 2.3)
    make_card(slide, left, top, Inches(3.8), Inches(2.0), title, [desc], color, title_size=15, item_size=12)

# ═══════════════════════════════════════════════════════
#  SLIDE 4 — SCOPE
# ═══════════════════════════════════════════════════════

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_shape_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT_CYAN)

add_section_number(slide, 2, ACCENT_CYAN)
add_slide_title(slide, "Project Scope", accent_color=ACCENT_CYAN)

# In-scope card
make_card(slide, Inches(0.6), Inches(2.3), Inches(5.8), Inches(4.6),
          "IN SCOPE",
          [
              "Physical Asset Management — books, journals, media",
              "Digital Resource Management — e-books, audiobooks",
              "Member Management — registration, profiling, tracking",
              "Circulation — issue, return, renewal workflows",
              "Reservation System — automated FIFO queuing",
              "Fine Management — auto-calculate & track fines",
              "Digital Access Logging — usage analytics",
          ], ACCENT_GREEN, title_size=16, item_size=13)

# Out of scope card
make_card(slide, Inches(6.8), Inches(2.3), Inches(5.8), Inches(4.6),
          "OUT OF SCOPE",
          [
              "Inter-library loan systems",
              "E-commerce / payment gateway integration",
              "AI-driven recommendation engines",
          ], ACCENT_ROSE, title_size=16, item_size=13)

# ═══════════════════════════════════════════════════════
#  SLIDE 5 — OBJECTIVES
# ═══════════════════════════════════════════════════════

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_shape_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT_CYAN)

add_section_number(slide, 2, ACCENT_CYAN)
add_slide_title(slide, "Objectives", accent_color=ACCENT_CYAN)

objectives = [
    ("Comprehensive Catalogue", "Centralize all physical & digital assets"),
    ("Accurate Circulation", "Real-time issue/return/inventory tracking"),
    ("Automated Fines", "Rule-based overdue fine calculation"),
    ("Efficient Reservations", "Fair, automated FIFO reservation queue"),
    ("Digital Access Logging", "Track e-book/journal usage & engagement"),
    ("Insightful Reports", "Overdue, popular books, fine collection reports"),
    ("Data Integrity", "PK, FK, CHECK constraints throughout"),
    ("Role-Based Access", "Admin, Librarian, Member permissions"),
    ("Fast Search", "Indexed multi-parameter catalogue search"),
    ("Streamlined Admin", "Reduce manual workload for staff"),
]

colors = [ACCENT_BLUE, ACCENT_CYAN, ACCENT_GREEN, ACCENT_PURPLE, ACCENT_AMBER,
          ACCENT_ROSE, ACCENT_BLUE, ACCENT_CYAN, ACCENT_GREEN, ACCENT_PURPLE]

for i, ((title, desc), color) in enumerate(zip(objectives, colors)):
    col = i % 5
    row = i // 5
    left = Inches(0.4 + col * 2.55)
    top = Inches(2.3 + row * 2.5)
    card = add_shape_rect(slide, left, top, Inches(2.35), Inches(2.1), BG_CARD, BORDER_COLOR, corner_radius=True)
    add_accent_bar(slide, left, top, Inches(2.35), Inches(0.04), color)

    # Number
    num_badge = add_shape_rect(slide, left + Inches(0.12), top + Inches(0.2), Inches(0.4), Inches(0.4), color, corner_radius=True)
    set_text(num_badge, str(i+1), font_size=14, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    add_text_box(slide, left + Inches(0.6), top + Inches(0.22), Inches(1.6), Inches(0.35),
                 title, font_size=12, color=WHITE, bold=True)
    add_text_box(slide, left + Inches(0.12), top + Inches(0.75), Inches(2.1), Inches(1.2),
                 desc, font_size=11, color=MID_GRAY)

# ═══════════════════════════════════════════════════════
#  SLIDE 6 — USERS AND USER ROLES
# ═══════════════════════════════════════════════════════

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_shape_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT_GREEN)

add_section_number(slide, 3, ACCENT_GREEN)
add_slide_title(slide, "Users & User Roles", accent_color=ACCENT_GREEN)

roles = [
    ("Admin / Chief Librarian", [
        "Full CRUD access to entire database",
        "Manage catalogue & member accounts",
        "Process issues/returns, adjust fines",
        "Override defaults, generate all reports",
    ], ACCENT_BLUE),
    ("Assistant Librarian", [
        "Process day-to-day issues & returns",
        "Search catalogue, manage queries",
        "Administer reservation queues",
        "Cannot modify system config",
    ], ACCENT_CYAN),
    ("Member", [
        "Search catalogue, view availability",
        "Place/cancel reservations",
        "View borrowing history & fines",
        "Access digital resources",
    ], ACCENT_GREEN),
    ("System / Automated", [
        "Daily fine calculation triggers",
        "Auto-expire unfulfilled reservations",
        "FIFO queue state management",
        "Overdue status updates",
    ], ACCENT_PURPLE),
]

for i, (role, perms, color) in enumerate(roles):
    left = Inches(0.5 + i * 3.15)
    make_card(slide, left, Inches(2.3), Inches(2.95), Inches(4.5),
              role, perms, color, title_size=14, item_size=12)

# ═══════════════════════════════════════════════════════
#  SLIDE 7 — FUNCTIONAL REQUIREMENTS (Part 1)
# ═══════════════════════════════════════════════════════

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_shape_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT_PURPLE)

add_section_number(slide, 4, ACCENT_PURPLE)
add_slide_title(slide, "Functional Requirements", "FR-001 through FR-012", accent_color=ACCENT_PURPLE)

fr_groups = [
    ("Catalogue Management", [
        "FR-001: Add/edit/remove books, authors, publishers, categories",
        "FR-002: Manage multiple editions and copies per title",
        "FR-003: Enforce unique ISBN and Accession Numbers",
    ], ACCENT_BLUE),
    ("Member Management", [
        "FR-004: Register new members, update profiles",
        "FR-005: Categorize by type (Student/Faculty/External)",
        "FR-006: Track membership validity periods",
    ], ACCENT_CYAN),
    ("Circulation (Issue/Return)", [
        "FR-007: Issue physical copies to valid members",
        "FR-008: One active issue per copy at a time",
        "FR-009: Record returns, calculate overdue days",
    ], ACCENT_GREEN),
    ("Reservation Management", [
        "FR-010: Place reservations on unavailable titles",
        "FR-011: Maintain FIFO reservation queue",
        "FR-012: Auto-expire unclaimed reservations",
    ], ACCENT_PURPLE),
]

for i, (title, items, color) in enumerate(fr_groups):
    col = i % 2
    row = i // 2
    left = Inches(0.5 + col * 6.2)
    top = Inches(2.5 + row * 2.4)
    make_card(slide, left, top, Inches(5.9), Inches(2.1), title, items, color, title_size=14, item_size=12)

# ═══════════════════════════════════════════════════════
#  SLIDE 8 — FUNCTIONAL REQUIREMENTS (Part 2)
# ═══════════════════════════════════════════════════════

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_shape_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT_PURPLE)

add_section_number(slide, 4, ACCENT_PURPLE)
add_slide_title(slide, "Functional Requirements (cont.)", "FR-013 through FR-019", accent_color=ACCENT_PURPLE)

fr_groups2 = [
    ("Fine Management", [
        "FR-013: Auto-calculate fines based on overdue days",
        "FR-014: Track partial and full fine payments",
        "FR-015: Fine amounts must never be negative",
    ], ACCENT_AMBER),
    ("Digital Resource Management", [
        "FR-016: Maintain records for e-books, journals, audiobooks",
        "FR-017: Log every member download/view access securely",
    ], ACCENT_ROSE),
    ("Search & Reports", [
        "FR-018: Search by title, author, ISBN, category",
        "FR-019: Reports: available titles, overdue copies,",
        "    member history, popular books, fine collection",
    ], ACCENT_BLUE),
]

for i, (title, items, color) in enumerate(fr_groups2):
    left = Inches(0.5 + i * 4.15)
    make_card(slide, left, Inches(2.5), Inches(3.9), Inches(2.5), title, items, color, title_size=14, item_size=12)

# ═══════════════════════════════════════════════════════
#  SLIDE 9 — NON-FUNCTIONAL REQUIREMENTS
# ═══════════════════════════════════════════════════════

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_shape_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT_AMBER)

add_section_number(slide, 5, ACCENT_AMBER)
add_slide_title(slide, "Non-Functional Requirements", accent_color=ACCENT_AMBER)

nfrs = [
    ("Data Integrity", "Normalization up to 3NF,\nconstraints to prevent\norphaned records", ACCENT_BLUE),
    ("Performance", "Efficient queries with\nproper indexing on\nISBN, Title, Author", ACCENT_CYAN),
    ("Security", "Role-Based Access Control\n(RBAC) for all user\ntypes", ACCENT_GREEN),
    ("Scalability", "Handle thousands of\nmembers and millions\nof catalogue items", ACCENT_PURPLE),
    ("Backup & Recovery", "Scheduled automated\nbackups with swift\nrecovery mechanisms", ACCENT_AMBER),
]

for i, (title, desc, color) in enumerate(nfrs):
    left = Inches(0.4 + i * 2.55)
    card = add_shape_rect(slide, left, Inches(2.5), Inches(2.35), Inches(3.0), BG_CARD, BORDER_COLOR, corner_radius=True)
    add_accent_bar(slide, left, Inches(2.5), Inches(2.35), Inches(0.05), color)

    # Icon-like colored circle
    icon = slide.shapes.add_shape(MSO_SHAPE.OVAL, left + Inches(0.8), Inches(2.8), Inches(0.6), Inches(0.6))
    icon.fill.solid()
    icon.fill.fore_color.rgb = color
    icon.line.fill.background()
    set_text(icon, str(i+1), font_size=18, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    add_text_box(slide, left + Inches(0.15), Inches(3.55), Inches(2.05), Inches(0.35),
                 title, font_size=14, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, left + Inches(0.15), Inches(3.95), Inches(2.05), Inches(1.4),
                 desc, font_size=11, color=MID_GRAY, alignment=PP_ALIGN.CENTER)

# ═══════════════════════════════════════════════════════
#  SLIDE 10 — BUSINESS RULES
# ═══════════════════════════════════════════════════════

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_shape_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT_AMBER)

add_section_number(slide, 5, ACCENT_AMBER)
add_slide_title(slide, "Business Rules", "Core constraints governing the database design", accent_color=ACCENT_AMBER)

rules = [
    ("BR-01", "ISBN must be unique across all books", ACCENT_BLUE),
    ("BR-02", "Accession number unique across all copies", ACCENT_CYAN),
    ("BR-03", "Only one active issue per copy at a time", ACCENT_GREEN),
    ("BR-04", "Issue_Date <= Due_Date always", ACCENT_PURPLE),
    ("BR-05", "Return_Date >= Issue_Date when returned", ACCENT_AMBER),
    ("BR-06", "Reservations follow FIFO queue order", ACCENT_ROSE),
    ("BR-07", "Fine amounts must be non-negative (>= 0)", ACCENT_BLUE),
    ("BR-08", "Max N active issues per member (by type)", ACCENT_CYAN),
    ("BR-09", "Digital logs: Member + Resource + Timestamp", ACCENT_GREEN),
    ("BR-10", "Every copy linked to a valid edition", ACCENT_PURPLE),
]

for i, (rule_id, desc, color) in enumerate(rules):
    col = i % 2
    row = i // 2
    left = Inches(0.6 + col * 6.2)
    top = Inches(2.5 + row * 0.95)

    # Rule badge
    badge = add_shape_rect(slide, left, top, Inches(0.85), Inches(0.7), color, corner_radius=True)
    set_text(badge, rule_id, font_size=11, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    # Rule card
    rule_card = add_shape_rect(slide, left + Inches(1.0), top, Inches(5.0), Inches(0.7), BG_CARD, BORDER_COLOR, corner_radius=True)
    set_text(rule_card, desc, font_size=13, color=LIGHT_GRAY)
    rule_card.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
    rule_card.text_frame.margin_left = Inches(0.15)

# ═══════════════════════════════════════════════════════
#  SLIDE 11 — ER DIAGRAM
# ═══════════════════════════════════════════════════════

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_shape_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT_ROSE)

add_section_number(slide, 6, ACCENT_ROSE)
add_slide_title(slide, "Entity-Relationship Diagram",
                "Conceptual design covering 14 entities with cardinality notation", accent_color=ACCENT_ROSE)

er_img_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "ER_Diagram.png")
if os.path.exists(er_img_path):
    # White background card for the image
    img_card = add_shape_rect(slide, Inches(0.8), Inches(2.3), Inches(11.7), Inches(4.8), RGBColor(0xFF, 0xFF, 0xFF), BORDER_COLOR, corner_radius=True)
    slide.shapes.add_picture(er_img_path, Inches(1.0), Inches(2.4), Inches(11.3), Inches(4.6))
else:
    add_text_box(slide, Inches(2), Inches(3.5), Inches(9), Inches(1),
                 "[ER_Diagram.png not found in project folder]", font_size=18, color=ACCENT_ROSE, alignment=PP_ALIGN.CENTER)

# ═══════════════════════════════════════════════════════
#  SLIDE 12 — ENTITIES SUMMARY
# ═══════════════════════════════════════════════════════

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_shape_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT_ROSE)

add_section_number(slide, 6, ACCENT_ROSE)
add_slide_title(slide, "Entities Overview", "14 tables covering all principal domains", accent_color=ACCENT_ROSE)

entities_grid = [
    ("Book", "ISBN, Title, Language", ACCENT_BLUE),
    ("Author", "Name, Bio", ACCENT_CYAN),
    ("Publisher", "Name, Address, Contact", ACCENT_GREEN),
    ("Category", "Name, Description", ACCENT_PURPLE),
    ("Edition", "EditionNo, Date, Format", ACCENT_AMBER),
    ("Copy", "AccessionNo, Status, Shelf", ACCENT_ROSE),
    ("Member", "Name, Email, Type, Active", ACCENT_BLUE),
    ("Issue", "IssueDate, DueDate, Return", ACCENT_CYAN),
    ("Reservation", "Date, Queue, Status", ACCENT_GREEN),
    ("Fine", "Amount, Reason, PaidDate", ACCENT_PURPLE),
    ("DigitalResource", "Type, Format, URL", ACCENT_AMBER),
    ("DigitalAccessLog", "AccessDate, Type, IP", ACCENT_ROSE),
    ("BookAuthor", "BookID, AuthorID (M:N)", MID_GRAY),
    ("BookCategory", "BookID, CategoryID (M:N)", MID_GRAY),
]

for i, (name, attrs, color) in enumerate(entities_grid):
    col = i % 7
    row = i // 7
    left = Inches(0.3 + col * 1.85)
    top = Inches(2.4 + row * 2.4)

    card = add_shape_rect(slide, left, top, Inches(1.7), Inches(2.0), BG_CARD, BORDER_COLOR, corner_radius=True)
    add_accent_bar(slide, left, top, Inches(1.7), Inches(0.04), color)

    add_text_box(slide, left + Inches(0.08), top + Inches(0.15), Inches(1.55), Inches(0.35),
                 name, font_size=12, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

    add_accent_bar(slide, left + Inches(0.15), top + Inches(0.55), Inches(1.4), Inches(0.01), BORDER_COLOR)

    add_text_box(slide, left + Inches(0.08), top + Inches(0.65), Inches(1.55), Inches(1.2),
                 attrs, font_size=10, color=MID_GRAY, alignment=PP_ALIGN.CENTER)

# ═══════════════════════════════════════════════════════
#  SLIDE 13 — RELATIONAL SCHEMA DIAGRAM
# ═══════════════════════════════════════════════════════

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_shape_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT_BLUE)

add_section_number(slide, 7, ACCENT_BLUE)
add_slide_title(slide, "Relational Schema Diagram",
                "Tables with data types, primary keys, and foreign key relationships", accent_color=ACCENT_BLUE)

schema_img_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Relational_Schema.png")
if os.path.exists(schema_img_path):
    img_card = add_shape_rect(slide, Inches(0.8), Inches(2.3), Inches(11.7), Inches(4.8), RGBColor(0xFF, 0xFF, 0xFF), BORDER_COLOR, corner_radius=True)
    slide.shapes.add_picture(schema_img_path, Inches(1.0), Inches(2.4), Inches(11.3), Inches(4.6))
else:
    add_text_box(slide, Inches(2), Inches(3.5), Inches(9), Inches(1),
                 "[Relational_Schema.png not found in project folder]", font_size=18, color=ACCENT_ROSE, alignment=PP_ALIGN.CENTER)

# ═══════════════════════════════════════════════════════
#  SLIDE 14 — KEY CONSTRAINTS
# ═══════════════════════════════════════════════════════

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_shape_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT_BLUE)

add_section_number(slide, 7, ACCENT_BLUE)
add_slide_title(slide, "Integrity Constraints Summary", accent_color=ACCENT_BLUE)

make_card(slide, Inches(0.5), Inches(2.3), Inches(3.8), Inches(4.6),
          "UNIQUE Constraints",
          ["Book.ISBN", "Copy.AccessionNumber", "Member.Email", "Publisher.Name", "Category.Name"],
          ACCENT_BLUE, title_size=14, item_size=12)

make_card(slide, Inches(4.6), Inches(2.3), Inches(3.8), Inches(4.6),
          "CHECK Constraints",
          ["Fine.Amount >= 0", "Book.TotalPages > 0", "Edition.EditionNumber > 0",
           "MaxBooksAllowed > 0", "QueuePosition >= 1",
           "DueDate >= IssueDate", "ReturnDate >= IssueDate"],
          ACCENT_CYAN, title_size=14, item_size=12)

make_card(slide, Inches(8.7), Inches(2.3), Inches(4.0), Inches(4.6),
          "DEFAULT Values",
          ["Copy.Status = 'Available'", "Issue.Status = 'Active'",
           "Reservation.Status = 'Pending'", "Fine.Status = 'Pending'",
           "Member.IsActive = TRUE"],
          ACCENT_GREEN, title_size=14, item_size=12)

# ═══════════════════════════════════════════════════════
#  SLIDE 15 — THANK YOU
# ═══════════════════════════════════════════════════════

slide = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(slide)
add_shape_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.06), ACCENT_BLUE)

# Decorative circles
circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-1), Inches(4), Inches(5), Inches(5))
circle.fill.solid()
circle.fill.fore_color.rgb = RGBColor(0x14, 0x17, 0x25)
circle.line.fill.background()

circle2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10), Inches(-1), Inches(4.5), Inches(4.5))
circle2.fill.solid()
circle2.fill.fore_color.rgb = RGBColor(0x1A, 0x1D, 0x2E)
circle2.line.fill.background()

add_text_box(slide, Inches(3), Inches(2.0), Inches(7), Inches(1.2),
             "Thank You", font_size=52, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)

add_accent_bar(slide, Inches(5.5), Inches(3.3), Inches(2.3), Inches(0.06), ACCENT_BLUE)

add_text_box(slide, Inches(3), Inches(3.7), Inches(7), Inches(0.5),
             "Library Circulation & Digital Resource Management System",
             font_size=16, color=MID_GRAY, alignment=PP_ALIGN.CENTER)

add_text_box(slide, Inches(3), Inches(4.4), Inches(7), Inches(0.5),
             "Review 1 Presentation  |  DBMS Project Based Learning",
             font_size=13, color=RGBColor(0x64, 0x74, 0x8B), alignment=PP_ALIGN.CENTER)

add_text_box(slide, Inches(3), Inches(5.3), Inches(7), Inches(0.4),
             "Questions & Discussion",
             font_size=18, color=ACCENT_CYAN, bold=True, alignment=PP_ALIGN.CENTER)

# ── Save ──
output = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Review1_Presentation.pptx")
prs.save(output)
print(f"Presentation saved: {output}")
