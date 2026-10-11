"""
Generate 5-Slide Presentation for DBMS PBL
Project: Library Circulation and Digital Resource Management System
STRICT REQUIREMENT: ZERO em dashes and ZERO en dashes.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

prs = Presentation()
prs.slide_width = Inches(13.333)  # 16:9 widescreen
prs.slide_height = Inches(7.5)

# Color Palette
COLOR_NAVY = RGBColor(0x0F, 0x4C, 0x81)
COLOR_TEAL = RGBColor(0x1B, 0x6C, 0xA8)
COLOR_DARK = RGBColor(0x2D, 0x37, 0x48)
COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
COLOR_MUTED = RGBColor(0x71, 0x80, 0x96)
COLOR_LIGHT_BG = RGBColor(0xF7, 0xFA, 0xFC)

def add_header(slide, title_text, category_text="DBMS COURSE PROJECT"):
    # Category Tracker
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.5), Inches(0.4))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.text = category_text.upper()
    p_c.font.size = Pt(11)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_TEAL

    # Main Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.9), Inches(11.5), Inches(0.8))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(26)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY

# ==============================================================================
# SLIDE 1: TITLE SLIDE
# ==============================================================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])

# Title text box
tb1 = slide1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(3.0))
tf1 = tb1.text_frame
tf1.word_wrap = True

p1 = tf1.paragraphs[0]
p1.text = "Library Circulation and Digital Resource\nManagement System"
p1.font.size = Pt(38)
p1.font.bold = True
p1.font.color.rgb = COLOR_NAVY

p2 = tf1.add_paragraph()
p2.text = "Database Management Systems: Project Based Learning (DBMS PBL)"
p2.font.size = Pt(18)
p2.font.color.rgb = COLOR_TEAL
p2.space_before = Pt(14)

p3 = tf1.add_paragraph()
p3.text = "Design, Normalization, Dual-Engine Relational Database, and Interactive Software Model"
p3.font.size = Pt(14)
p3.font.color.rgb = COLOR_MUTED
p3.space_before = Pt(8)

# Footer metadata
tb_foot = slide1.shapes.add_textbox(Inches(1.0), Inches(5.8), Inches(11.3), Inches(1.0))
tf_foot = tb_foot.text_frame
p_f = tf_foot.paragraphs[0]
p_f.text = "Department of Computer Science and Engineering | Academic Year 2025-2026"
p_f.font.size = Pt(12)
p_f.font.bold = True
p_f.font.color.rgb = COLOR_DARK

# ==============================================================================
# SLIDE 2: PROBLEM STATEMENT & OBJECTIVES
# ==============================================================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide2, "Problem Statement and Project Objectives")

# Left Column: Problems
tb_l2 = slide2.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(5.6), Inches(4.8))
tf_l2 = tb_l2.text_frame
tf_l2.word_wrap = True

p_lh = tf_l2.paragraphs[0]
p_lh.text = "Challenges in Legacy Library Systems:"
p_lh.font.size = Pt(18)
p_lh.font.bold = True
p_lh.font.color.rgb = COLOR_NAVY

bullet_probs = [
    ("Unreliable Asset Tracking: ", "Manual registers cause barcode discrepancies and misplaced physical copies."),
    ("No Digital Integration: ", "Older systems cannot monitor electronic downloads or manage e-books and audiobooks."),
    ("Manual Fine Errors: ", "Manual overdue calculations cause disputes, errors, and revenue leakage."),
    ("Unfair Waitlists: ", "Lack of automated First-In-First-Out queuing causes chaotic book reservations.")
]
for b_t, b_d in bullet_probs:
    p_b = tf_l2.add_paragraph()
    p_b.space_before = Pt(10)
    r1 = p_b.add_run()
    r1.text = b_t
    r1.font.bold = True
    r1.font.size = Pt(13)
    r1.font.color.rgb = COLOR_DARK
    r2 = p_b.add_run()
    r2.text = b_d
    r2.font.size = Pt(13)
    r2.font.color.rgb = COLOR_MUTED

# Right Column: Objectives
tb_r2 = slide2.shapes.add_textbox(Inches(6.8), Inches(1.9), Inches(5.6), Inches(4.8))
tf_r2 = tb_r2.text_frame
tf_r2.word_wrap = True

p_rh = tf_r2.paragraphs[0]
p_rh.text = "Target System Objectives:"
p_rh.font.size = Pt(18)
p_rh.font.bold = True
p_rh.font.color.rgb = COLOR_TEAL

bullet_objs = [
    ("Unified Centralized Catalog: ", "Manage physical editions, authors, publishers, and digital assets."),
    ("Automated Overdue Fines: ", "Enforce mathematical penalty rules with automated check-in assessments."),
    ("Strict FIFO Reservation: ", "Guarantee ordered book waitlists with automated queue allocation."),
    ("Security & Audit Logging: ", "Maintain immutable patron access logs for all electronic media downloads.")
]
for b_t, b_d in bullet_objs:
    p_b = tf_r2.add_paragraph()
    p_b.space_before = Pt(10)
    r1 = p_b.add_run()
    r1.text = b_t
    r1.font.bold = True
    r1.font.size = Pt(13)
    r1.font.color.rgb = COLOR_DARK
    r2 = p_b.add_run()
    r2.text = b_d
    r2.font.size = Pt(13)
    r2.font.color.rgb = COLOR_MUTED

# ==============================================================================
# SLIDE 3: DATABASE ARCHITECTURE & RELATIONAL SCHEMA
# ==============================================================================
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide3, "Relational Schema and 3NF Normalization")

tb3 = slide3.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(11.6), Inches(5.0))
tf3 = tb3.text_frame
tf3.word_wrap = True

p3_top = tf3.paragraphs[0]
p3_top.text = "Normalized Relational Architecture (14 Cohesive Tables):"
p3_top.font.size = Pt(18)
p3_top.font.bold = True
p3_top.font.color.rgb = COLOR_NAVY

schema_groups = [
    ("Catalog & Inventory Entities: ", "Book, Author, BookAuthor, Publisher, Category, BookCategory, Edition, Copy"),
    ("Patron & Access Management: ", "Member (Students, Faculty, Staff, Public), Role-based borrowing quotas"),
    ("Circulation Operations: ", "Issue (Checkouts and returns), Fine (Automated overdue penalties)"),
    ("Resource Availability: ", "Reservation (FIFO waitlist queue with priority ordering)"),
    ("Electronic Media Management: ", "DigitalResource (PDF, EPUB, MP3), DigitalAccessLog (Audit trail with IP and timestamp)")
]

for g_t, g_d in schema_groups:
    p_g = tf3.add_paragraph()
    p_g.space_before = Pt(10)
    r1 = p_g.add_run()
    r1.text = g_t
    r1.font.bold = True
    r1.font.size = Pt(14)
    r1.font.color.rgb = COLOR_TEAL
    r2 = p_g.add_run()
    r2.text = g_d
    r2.font.size = Pt(14)
    r2.font.color.rgb = COLOR_DARK

p3_note = tf3.add_paragraph()
p3_note.space_before = Pt(16)
p3_note.text = "Normalization Guarantee: Verified in 3NF and BCNF to ensure zero insertion, update, or deletion anomalies."
p3_note.font.size = Pt(13)
p3_note.font.bold = True
p3_note.font.color.rgb = COLOR_NAVY

# ==============================================================================
# SLIDE 4: WORKING MODEL IMPLEMENTATION & FEATURES
# ==============================================================================
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide4, "Working Model Implementation and Features")

# Left Box: Tech Stack
tb_l4 = slide4.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(5.6), Inches(4.8))
tf_l4 = tb_l4.text_frame
tf_l4.word_wrap = True

p_l4h = tf_l4.paragraphs[0]
p_l4h.text = "System Architecture and Technology Stack:"
p_l4h.font.size = Pt(18)
p_l4h.font.bold = True
p_l4h.font.color.rgb = COLOR_NAVY

stack_items = [
    ("UI & Dashboard: ", "Streamlit web application with custom CSS and navigation"),
    ("Analytics & Visuals: ", "Plotly charts for inventory status, member categories, and trends"),
    ("Persistence Layer: ", "Dual-engine architecture: MySQL 8.4 Server with SQLite auto-failover"),
    ("Data Operations: ", "Pandas 2.3 for transactional dataframes and complex relational queries")
]
for s_t, s_d in stack_items:
    p_s = tf_l4.add_paragraph()
    p_s.space_before = Pt(10)
    r1 = p_s.add_run()
    r1.text = s_t
    r1.font.bold = True
    r1.font.size = Pt(13)
    r1.font.color.rgb = COLOR_DARK
    r2 = p_s.add_run()
    r2.text = s_d
    r2.font.size = Pt(13)
    r2.font.color.rgb = COLOR_MUTED

# Right Box: Functional Modules
tb_r4 = slide4.shapes.add_textbox(Inches(6.8), Inches(1.9), Inches(5.6), Inches(4.8))
tf_r4 = tb_r4.text_frame
tf_r4.word_wrap = True

p_r4h = tf_r4.paragraphs[0]
p_r4h.text = "Implemented Core Modules:"
p_r4h.font.size = Pt(18)
p_r4h.font.bold = True
p_r4h.font.color.rgb = COLOR_TEAL

modules = [
    ("Live Circulation Desk: ", "Enforces member borrowing limits; performs 1-click checkouts and returns."),
    ("Automated Overdue Fines: ", "Dynamically computes days past due date and assesses late penalties."),
    ("Reservation Waitlist: ", "Tracks FIFO waitlist queue and assigns priority positions."),
    ("Interactive SQL Runner: ", "Executes multi-table joins, subqueries, and aggregate evaluations in real time.")
]
for m_t, m_d in modules:
    p_m = tf_r4.add_paragraph()
    p_m.space_before = Pt(10)
    r1 = p_m.add_run()
    r1.text = m_t
    r1.font.bold = True
    r1.font.size = Pt(13)
    r1.font.color.rgb = COLOR_DARK
    r2 = p_m.add_run()
    r2.text = m_d
    r2.font.size = Pt(13)
    r2.font.color.rgb = COLOR_MUTED

# ==============================================================================
# SLIDE 5: RESULTS, TESTING & CONCLUSION
# ==============================================================================
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
add_header(slide5, "Results, Verification Testing, and Conclusion")

tb5 = slide5.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(11.6), Inches(5.0))
tf5 = tb5.text_frame
tf5.word_wrap = True

p5_top = tf5.paragraphs[0]
p5_top.text = "Verification Results and Project Summary:"
p5_top.font.size = Pt(18)
p5_top.font.bold = True
p5_top.font.color.rgb = COLOR_NAVY

results_items = [
    ("Constraint Enforcement: ", "Verified 100% rejection of duplicate ISBNs and barcodes via unique constraints."),
    ("Quota Control: ", "Successfully blocked checkout attempts once patron max allowed book limits were reached."),
    ("Fine Accuracy: ", "Calculated overdue fines with zero computational errors across all return scenarios."),
    ("Audit Trail Integrity: ", "Every digital download accurately captured patron ID, timestamp, and client IP address."),
    ("Demonstrated Impact: ", "Provides an end-to-end, error-free management system for modern hybrid libraries.")
]

for r_t, r_d in results_items:
    p_r = tf5.add_paragraph()
    p_r.space_before = Pt(10)
    r1 = p_r.add_run()
    r1.text = r_t
    r1.font.bold = True
    r1.font.size = Pt(14)
    r1.font.color.rgb = COLOR_TEAL
    r2 = p_r.add_run()
    r2.text = r_d
    r2.font.size = Pt(14)
    r2.font.color.rgb = COLOR_DARK

p5_concl = tf5.add_paragraph()
p5_concl.space_before = Pt(18)
p5_concl.text = "Conclusion: The project achieves all DBMS PBL objectives with a normalized database and a fully functional working model."
p5_concl.font.size = Pt(14)
p5_concl.font.bold = True
p5_concl.font.color.rgb = COLOR_NAVY

# Save presentation
out_pptx = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'Final_Presentation_5Slides.pptx')
prs.save(out_pptx)
print("Presentation generated successfully at:", out_pptx)
